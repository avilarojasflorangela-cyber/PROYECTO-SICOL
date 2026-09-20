from Model_Usuarios import Model_Usuarios
from Model_Proveedores import Model_Proveedores
from Model_OrdenCompra import Model_OrdenCompra
from Model_Inventario import Model_Inventario
from Model_Facturas import Model_Facturas
from Model_EntradaMaterial import Model_EntradaMaterial
from Model_Empleados import Model_Empleados
from Model_DocumentosProveedor import Model_DocumentosProveedor
from Model_DetalleProducto import Model_DetalleProducto
from Model_CrearPedido import Model_CrearPedido
from Model_ConsultarEstado import Model_ConsultarEstado
from Model_HistorialEstado import Model_HistorialEstado
from UserLogin import User
from openpyxl import Workbook
from io import BytesIO
from flask import Flask, jsonify, request, send_file
from datetime import datetime, timedelta
import os
from flask_cors import CORS

app = Flask(__name__)
app.config ['SECRET_KEY'] = '7110c8ae51a4b5af97be6534caef90e4bb9bdcb3380af008f90b23a5d1616bf319bc298105da20fe'

CORS(app)
usuarios = Model_Usuarios()
proveedores = Model_Proveedores()
orden_compra = Model_OrdenCompra()
inventario = Model_Inventario()
facturas = Model_Facturas()
entrada_material = Model_EntradaMaterial()
empleados = Model_Empleados()
documentos_proveedor = Model_DocumentosProveedor()
detalle_producto = Model_DetalleProducto()
crear_pedido =  Model_CrearPedido()
consultar_estado = Model_ConsultarEstado()
historial_estado = Model_HistorialEstado()

@app.route('/login', methods=['POST'])
def loginUsuarios():

    ID_CC = request.json['ID_CC']
    password = request.json['password']
    resultado = usuarios.verificarUsuario(ID_CC)
    if resultado:
        usuario = User(
            resultado[0],
            resultado[1],
            resultado[2],
            resultado[3],
            resultado[4],
            resultado[5],
            resultado[6],
        )
        if usuario.check_password(password):
            return jsonify({
                "mensaje": "Inicio de sesión correcto",
                "id": usuario.id,
                "ID_CC": usuario.ID_CC,
                "nombre": usuario.nombre,
                "apellido": usuario.apellido,
                "correo": usuario.correo,
                "area": usuario.area
            })
    return jsonify({
        "mensaje": "Contraseña o usuario incorrecto"
    }), 401

@app.route('/usuarios', methods=['GET'])
def listar_usuarios():
    return jsonify(usuarios.obtener_usuarios())

@app.route('/usuario/<int:id>', methods=['GET'])
def listar_usuario(id):
    return jsonify(usuarios.obtener_usuario(id))

@app.route('/nuevo_usuario', methods=['POST'])
def crear_usuario():
    ID_CC = request.json['ID_CC']
    Usua_Nombre = request.json['Usua_Nombre']
    Usua_Apellido = request.json['Usua_Apellido']
    Usua_Ciudad = request.json['Usua_Ciudad']
    Usua_Telefono = request.json['Usua_Telefono']
    Usua_Correo = request.json['Usua_Correo']
    Usua_Estado = request.json['Usua_Estado']
    Usua_Area = request.json['Usua_Area']
    contrasena = request.json['contrasena']

    usuarios.insertar_usuario(ID_CC, Usua_Nombre, Usua_Apellido, Usua_Ciudad, Usua_Telefono, Usua_Correo, Usua_Estado, Usua_Area, contrasena)
    return jsonify({"mensaje": "Usuario creado correctamente."})

@app.route('/actualizar_usuario/<int:id>', methods=['PUT'])
def actualizar_usuario(id):
    data = request.json
    ID_CC = data.get('ID_CC')
    Usua_Nombre = data.get('Usua_Nombre')
    Usua_Apellido = data.get('Usua_Apellido')
    Usua_Ciudad = data.get('Usua_Ciudad')
    Usua_Telefono = data.get('Usua_Telefono')
    Usua_Correo = data.get('Usua_Correo')
    Usua_Estado = data.get('Usua_Estado')
    Usua_Area = data.get('Usua_Area')
    contrasena = data.get('contrasena')

    if ID_CC is None and Usua_Nombre is None and Usua_Apellido is None and Usua_Ciudad is None and Usua_Telefono is None and Usua_Correo is None and Usua_Estado is None and Usua_Area is None and contrasena is None:
        return jsonify({"error": "No se pudo actualizar datos."}), 400

    usuarios_existente = usuarios.obtener_usuario(id)

    if usuarios_existente:
        usuarios.actualizar_usuario(id, ID_CC, Usua_Nombre, Usua_Apellido, Usua_Ciudad, Usua_Telefono, Usua_Correo, Usua_Estado, Usua_Area, contrasena)
        return jsonify({"mensaje": "usuarios actualizado correctamente."})
    else:
        return jsonify({"error": "El usuario no existe"}), 404

@app.route('/eliminar_usuario/<int:id>', methods=['DELETE'])
def eliminar_usuario(id):
    if usuarios.obtener_usuario(id):
        usuarios.eliminar_usuario(id)
        return jsonify({"mensaje": "Usuario eliminado correctamente."})
    else:
        return jsonify({"error": "El usuario no existe."}), 404

@app.route('/proveedores', methods=['GET'])
def listar_proveedores():
    return jsonify(proveedores.obtener_proveedores())

@app.route('/proveedor/<int:id>', methods=['GET'])
def listar_proveedor(id):
    return jsonify(proveedores.obtener_proveedor(id))

@app.route('/nuevo_proveedor', methods=['POST'])
def crear_proveedor():
    ID_CC_NIT = request.json['ID_CC_NIT']
    razon_social = request.json['razon_social']
    nombre_asesor = request.json['nombre_asesor']
    correo_electronico = request.json['correo_electronico']
    numero_telefono = request.json['numero_telefono']
    ciudad = request.json['ciudad']
    pais = request.json['pais']
    forma_pago = request.json['forma_pago']
    estado = 'Pendiente'
    id_usuario = request.json['id_usuario']

    proveedores.insertar_proveedor(ID_CC_NIT, razon_social, nombre_asesor,correo_electronico, numero_telefono, ciudad, pais, forma_pago, estado, id_usuario)
    return jsonify({"mensaje": "Proveedor creado correctamente"})

@app.route('/actualizar_proveedor/<int:id>', methods=['PUT'])
def actualizar_proveedor(id):
    data = request.json
    ID_CC_NIT = data.get('ID_CC_NIT')
    razon_social = data.get('razon_social')
    nombre_asesor = data.get('nombre_asesor')
    correo_electronico = data.get('correo_electronico')
    numero_telefono = data.get('numero_telefono')
    ciudad = data.get('ciudad')
    pais = data.get('pais')
    forma_pago = data.get('forma_pago')
    estado = data.get('estado')
    id_usuario = data.get('id_usuario')

    if ID_CC_NIT is None and razon_social is None and nombre_asesor is None and correo_electronico is None and  numero_telefono is None and ciudad is None and pais is None and forma_pago is None and estado is None and id_usuario is None:
        return jsonify({"error": "No se pudo actualizar datos."}), 400
    
    proveedor_existente = proveedores.obtener_proveedor(id)

    if proveedor_existente:
        proveedores.actualizar_proveedor(id, ID_CC_NIT, razon_social, nombre_asesor, correo_electronico, numero_telefono, ciudad, pais, forma_pago, estado, id_usuario)
        return jsonify({"mensaje": "Proveedor actualizado correctamente"})
    else:
        return jsonify({"error": "El proveedor no existe"}), 404

@app.route('/actualizar_estado_proveedor/<int:id>', methods=['PUT'])
def actualizar_estado_proveedor(id):
    data = request.json
    estado = data.get('estado')
    proveedor_existente = proveedores.obtener_proveedor(id)

    if proveedor_existente:
        proveedores.actualizar_estado_proveedor(id, estado)
        return jsonify({"mensaje": "Estado del proveedor actualizado correctamente"})
    else:
        return jsonify({"error": "El proveedor no existe"}), 404

@app.route('/eliminar_proveedor/<int:id>', methods=['DELETE'])
def eliminar_proveedor(id):
    if proveedores.obtener_proveedor(id):
        proveedores.eliminar_proveedor(id)
        return jsonify({"mensaje": "Proveedor eliminado correctamente."})
    else:
        return jsonify({"error": "El proveedor no existe."}), 404

@app.route('/reporte_proveedores', methods=['GET'])
def reporte_proveedores():

    resultado = proveedores.reporte_proveedores()

    return jsonify(resultado)

@app.route('/reporte_proveedores_excel', methods=['GET'])
def reporte_proveedores_excel():

    datos = proveedores.reporte_proveedores()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Proveedores"

    encabezados = [
        "Id proveedor",
        "NIT",
        "Razon social",
        "Nombre del asesor",
        "Correo electronico",
        "Número de teléfono",
        "Ciudad",
        "País",
        "Forma de pago",
        "estado",
        "id_usuario"
    ]

    hoja.append(encabezados)

    for fila in datos:
        hoja.append(fila)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return send_file(
        archivo,
        as_attachment=True,
        download_name="reporte_proveedores.xlsx",
        mimetype="application/vnd.openxmlformats.officedocument.spreadsheetml.sheet"
    )

@app.route('/orden_compras', methods=['GET'])
def listar_orden_compras():
    return jsonify(orden_compra.obtener_orden_compras())

@app.route('/orden_compra/<int:id>', methods=['GET'])
def listar_orden_compra(id):
    return jsonify(orden_compra.obtener_orden_compra(id))

@app.route('/pedidos_aprobados_orden', methods=['GET'])
def listar_pedidos_aprobados_orden():
    return jsonify(orden_compra.obtener_pedidos_aprobados())

@app.route('/nueva_orden_compra', methods=['POST'])
def crear_orden_compra():
    precio_unitario = request.json['precio_unitario']
    cantidad = request.json['cantidad']
    fecha_orden = request.json['fecha_orden']
    fecha_entrega = request.json['fecha_entrega']
    estado_orden = 'Orden creada'
    observaciones = request.json['observaciones']
    id_proveedor = request.json['id_proveedor']
    proveedor = proveedores.obtener_proveedor(id_proveedor)
    forma_pago = proveedor[8]
    cod_producto = request.json['cod_producto']
    id_pedido = request.json['id_pedido']

    id_orden = orden_compra.insertar_orden_compra(precio_unitario, cantidad, fecha_orden, fecha_entrega, estado_orden, observaciones, id_proveedor, cod_producto, id_pedido)
    
    fecha_radicacion = datetime.now().date()
    fecha_vencimiento = fecha_radicacion + timedelta(days=10)
    total_factura = precio_unitario * cantidad
    estado_factura = 'Pendiente de pago'

    facturas.insertar_factura(fecha_radicacion, fecha_vencimiento, forma_pago, total_factura, estado_factura, observaciones, id_proveedor, id_orden)
    return jsonify({"mensaje": "Orden de compra creada exitosamente"})

@app.route('/nueva_orden_compra/<int:id_pedido>', methods=['GET'])
def obtener_pedido_para_orden_compra(id_pedido):
    pedido = orden_compra.obtener_pedido_para_orden(id_pedido)
    if pedido:
        return jsonify(pedido) 
    else:
        return jsonify({"error": "El pedido no existe."}), 404

@app.route('/actualizar_orden_compra/<int:id>', methods=['PUT'])
def actualizar_orden_compra(id):
    data = request.json
    precio_unitario = data.get('precio_unitario')
    cantidad = data.get('cantidad')
    fecha_orden = data.get('fecha_orden')
    fecha_entrega = data.get('fecha_entrega')
    estado_orden = data.get('estado_orden')
    observaciones = data.get('observaciones')
    id_proveedor = data.get('id_proveedor') 
    cod_producto = data.get('cod_producto')
    id_pedido = data.get('id_pedido')

    if precio_unitario is None and cantidad is None and fecha_orden is None and fecha_entrega is None and estado_orden is None and observaciones is None and id_proveedor is None and cod_producto is None and id_pedido is None:
        return jsonify({"error": "No se pudo actualizar los datos."})

    orden_compra_existente = orden_compra.obtener_orden_compra(id)

    if orden_compra_existente:
        orden_compra.actualizar_orden_compra(id, precio_unitario, cantidad, fecha_orden, fecha_entrega, estado_orden, observaciones, id_proveedor, cod_producto, id_pedido)
        return jsonify({"mensaje": "Orden de compra actualizada correctamente."})
    else:
        return jsonify({"error": "La orden de compra no existe."}), 404

@app.route('/eliminar_orden_compra/<int:id>', methods=['DELETE'])
def eliminar_orden_compra(id):
    if orden_compra.obtener_orden_compra(id):
        orden_compra.eliminar_orden_compra(id)
        return jsonify({"mensaje": "Orden de compra eliminada correctamente."})
    else:
        return jsonify({"error": "La orden de compra no existe."}), 404

@app.route('/reporte_orden_compra', methods=['GET'])
def reporte_orden_compra():
    resultado = orden_compra.reporte_orden_compra()
    return jsonify(resultado)

@app.route('/reporte_orden_compra_excel', methods=['GET'])
def reporte_orden_compra_excel():
    datos = orden_compra.reporte_orden_compra()

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Ordenes de compras'

    encabezados = [
        "Id orden", 
        "Precio unitario",
        "Cantidad",
        "Fecha de orden", 
        "Fecha de entrega",
        "Forma de pago",
        "Estado de la orden",
        "Observaciones",
        "Nit proveedor",
        "Razon social",
        "Correo electronico",
        "Id pedido",
        "Nombre del producto",
        "Destino",
        "Fecha requerida",
        "Código del producto",
        "Descrición",
        "Unidad de medida",
        "Marca",
        "Categoria"
    ]

    hoja.append(encabezados)
    for fila in datos:
        hoja.append(fila)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return send_file(
        archivo,
        as_attachment=True,
        download_name="reporte_orden_compra.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route('/inventarios', methods=['GET'])
def listar_inventarios():
    return jsonify(inventario.obtener_inventarios())

@app.route('/inventario/<int:id>', methods=['GET'])
def listar_inventario(id):
    return jsonify(inventario.obtener_inventario(id))

@app.route('/inventario', methods=['GET'])
def listar_inventarios_join():
    return jsonify(inventario.inventario())


@app.route('/nuevo_inventario', methods=['POST'])
def crear_inventario():
    nombre_producto = request.json['nombre_producto']
    cantidad = request.json['cantidad']
    fecha_ingeso_material = request.json['fecha_ingreso_material']
    tipo_movimiento = request.json['tipo_movimiento']
    stock_minimo = request.json['stock_minimo']
    id_orden = request.json['id_orden']
    cod_producto = request.json['cod_producto']

    inventario.insertar_inventario(nombre_producto, cantidad, fecha_ingeso_material, tipo_movimiento, stock_minimo, id_orden, cod_producto)
    return jsonify({"mensaje": "Inventario creado correctamente."})

@app.route('/actualizar_inventario/<int:id>', methods=['PUT'])
def actualizar_inventario(id):
    data = request.json
    nombre_producto = data.get('nombre_producto')
    cantidad = data.get('cantidad')
    fecha_ingreso_material = data.get('fecha_ingreso_material')
    tipo_movimiento = data.get('tipo_movimiento')
    stock_minimo = data.get('stock_minimo')
    id_orden = data.get('id_orden')
    cod_producto = data.get('cod_producto')

    if nombre_producto is None and cantidad is None and fecha_ingreso_material is None and tipo_movimiento is None and stock_minimo is None and id_orden is None and cod_producto is None:
        return jsonify({"error": "No se pudo actualizar datos"}), 400
    inventario_existente = inventario.obtener_inventario(id)
    if inventario_existente:
        inventario.actualizar_inventario(id, nombre_producto, cantidad, fecha_ingreso_material, tipo_movimiento, stock_minimo, id_orden, cod_producto)
        return jsonify({"mensaje": "Inventario actualizado correctamente."})
    else:
        return jsonify({"error": "El inventario no existe."}), 404

@app.route('/eliminar_inventario/<int:id>', methods=['DELETE'])
def eliminar_inventario(id):
    if inventario.obtener_inventario(id):
        inventario.eliminar_inventario(id)
        return jsonify({"mensaje": "Inventario eliminado correctamente"})
    else:
        return jsonify({"error": "El inventario no existe"}), 404

@app.route('/sacar_inventario/<int:cod_producto>', methods=['PUT'])
def sacar_inventario(cod_producto):

    data = request.json
    cantidad_solicitada = data.get('cantidad')

    if cantidad_solicitada is None:
        return jsonify({"error": "Debe indicar la cantidad a retirar"}), 400
    
    if cantidad_solicitada <= 0:
        return jsonify({"error": "La cantidad debe ser mayor a cero"}),400
    
    producto = inventario.obtener_por_producto(cod_producto)

    if not producto:
        return jsonify({"error": "El producto no existe en el inventario"}), 400
    
    cantidad_actual = producto[2]
    stock_minimo = producto[5]
    cantidad_restante = cantidad_actual - cantidad_solicitada

    if cantidad_restante < stock_minimo:
        return jsonify({"error": "No se puede retirar esa cantidad porque el inventario quedaría por debajo del stock minimo", "cantidad_actual": cantidad_actual, "stock_minimo": stock_minimo}), 400
    inventario.actualizar_cantidad(cod_producto, cantidad_restante )

    return jsonify({"mensaje": "Producto retirado correctamente.", "cantidad_restante": cantidad_restante})

@app.route('/reporte_inventario', methods=['GET'])
def reporte_inventario():

    resultado = inventario.reporte_invetario()

    return jsonify(resultado)

@app.route('/reporte_inventario_excel', methods=['GET'])
def reporte_inventario_excel():

    datos = inventario.reporte_invetario()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Inventario"

    encabezados = [
        "Id inventario",
        "Nombre del producto",
        "Cantidad",
        "Fecha de ingreso",
        "Tipo de movimiento",
        "Stock mínimo",
        "Id_orden",
        "Razon social",
        "Código del producto",
        "Precio unitario",
        "Categoría",
        "Marca"
    ]

    hoja.append(encabezados)

    for fila in datos:
        hoja.append(fila)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return send_file(
        archivo,
        as_attachment=True,
        download_name="reporte_inventario.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    
@app.route('/facturas_total', methods=['GET'])
def listar_facturas_todas():
    return jsonify(facturas.obtener_facturas())

@app.route('/factura/<int:id>', methods=['GET'])
def listar_factura(id):
    return jsonify(facturas.obtener_factura(id))

@app.route('/facturas', methods=['GET'])
def listar_facturas():
    return jsonify(facturas.facturas())

@app.route('/nueva_factura', methods=['POST'])
def crear_factura():
    fecha_radicacion = request.json['fecha_radicacion']
    fecha_vencimiento = request.json['fecha_vencimiento']
    forma_pago = request.json['forma_pago']
    total_factura = request.json['total_factura']
    estado_factura = 'Pendiente de pago'
    observaciones = request.json['observaciones']
    id_proveedor = request.json['id_proveedor']
    id_orden = request.json['id_orden']

    facturas.insertar_factura(fecha_radicacion, fecha_vencimiento, forma_pago, total_factura, estado_factura, observaciones, id_proveedor, id_orden)
    return jsonify({"mensaje": "Factura creada correctamente."})

@app.route('/actualizar_factura/<int:id>', methods=['PUT'])
def actualizar_factura(id):
    data = request.json
    fecha_radicacion = data.get('fecha_radicacion')
    fecha_vencimiento = data.get('fecha_vencimiento')
    forma_pago = data.get('forma_pago')
    total_factura = data.get('total_factura')
    estado_factura = data.get('estado_factura')
    observaciones = data.get('observaciones')
    id_proveedor = data.get('id_proveedor')
    id_orden = data.get('id_orden')

    if fecha_radicacion is None and fecha_vencimiento is None and forma_pago is None and total_factura is None and estado_factura is None and observaciones is None and id_proveedor is None and id_orden is None:
        return jsonify({"error": "No se pudo actualizar datos."}), 400
    
    factura_existente = facturas.obtener_factura(id)

    if factura_existente:
        facturas.actualizar_factura(id, fecha_radicacion, fecha_vencimiento, forma_pago, total_factura, estado_factura, observaciones, id_proveedor, id_orden)
        return jsonify({"mensaje": "Factura actualizada correctamente."})
    else:
        return jsonify({"error": "La factura no existe."}), 400

@app.route('/actualizar_estado_factura/<int:id>', methods=['PUT'])
def actualizar_estado_factura(id):
    data = request.json
    estado_factura = data.get('estado_factura')
    if estado_factura is None:
        return jsonify({"error": "Debe indicar el estado de la factura."}), 400
    factura_existente = facturas.obtener_factura(id)
    if factura_existente:
        facturas.actualizar_estado_factura(id, estado_factura)
        return jsonify({"mensaje": "Estado de factura actualizado correctamente."})
    else:
        return jsonify({"error": "La factura no existe."}), 404

@app.route('/actualizar_observacion_factura/<int:id>', methods=['PUT'])
def actualizar_observacion_factura(id):
    data = request.json
    observaciones = data.get('observaciones')

    factura_existente = facturas.obtener_factura(id)

    if factura_existente:
        facturas.actualizar_observacion_factura(id, observaciones)
        return jsonify({"mensaje": "Observación actualizada correctamente."})
    else:
        return jsonify({"error": "La factura no existe."}), 404

@app.route('/eliminar_factura/<int:id>', methods=['DELETE'])
def eliminar_factura(id):
    if facturas.obtener_factura(id):
        facturas.eliminar_factura(id)
        return jsonify({"mensaje": "Factura eliminada correctamente."})
    else:
        return jsonify({"error": "La factura no existe."}), 404

@app.route('/reporte_facturas', methods=['GET'])
def reporte_facturas():

    resultado = facturas.reporte_facturas()

    return jsonify(resultado)

@app.route('/reporte_facturas_excel', methods=['GET'])
def reporte_facturas_excel():

    datos = facturas.reporte_facturas()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Facturas"

    encabezados = [
        "Id facturas",
        "Fecha de radicación",
        "Fecha de vencimiento",
        "Forma de pago",
        "Total de factura",
        "Estado de la factura",
        "Observaciones",
        "Nit",
        "Razon social",
        "Nombre del asesor",
        "Correo_electronico",
        "Número de teléfono",
        "Ciudad",
        "Id orden"
    ]

    hoja.append(encabezados)

    for fila in datos:
        hoja.append(fila)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return send_file(
        archivo,
        as_attachment=True,
        download_name="reporte_facturas.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route('/entrada_materiales', methods=['GET'])
def listar_entrada_materiales():
    return jsonify(entrada_material.obtener_entrada_materiales())

@app.route('/entrada_material/<int:id>', methods=['GET'])
def listar_entrada_material(id):
    return jsonify(entrada_material.obtener_entrada_material(id))

@app.route('/entrada_material', methods=['GET'])
def listar_entrada_material_join():
    return jsonify(entrada_material.entrada_material())

@app.route('/nueva_entrada_material', methods=['POST'])
def crear_entrada_material():
    fecha_entrada = request.json['fecha_entrada']
    cantidad_recibida = int(request.json['cantidad_recibida'])
    recibido_por = request.json['recibido_por']
    observaciones = request.json['observaciones']
    id_orden = int(request.json['id_orden'])
    cod_producto = int(request.json['cod_producto'])

    orden = orden_compra.obtener_orden_compra(id_orden)

    if not orden:
        return jsonify({"error": "La orden de compra no existe."}), 404

    cantidad_ordenada = orden[2]
    entradas = entrada_material.obtener_entradas_por_orden(id_orden)
    cantidad_anterior = 0

    for entrada in entradas:
        cantidad_anterior += int(entrada[2])

    cantidad_total = cantidad_anterior + cantidad_recibida

    if cantidad_total > cantidad_ordenada:
        return jsonify({"error": "La cantidad recibida supera la cantidad de la orden de compra."}), 400
    if cantidad_total == cantidad_ordenada:
        cumplimiento_orden_compra = "Cumplida"
    else:
        cumplimiento_orden_compra = "Parcial"

    entrada_material.insertar_entrada_material(fecha_entrada, cantidad_recibida, recibido_por, cumplimiento_orden_compra, observaciones, id_orden, cod_producto)
    producto_inventario = inventario.obtener_por_producto(cod_producto)

    if producto_inventario:
        cantidad_actual = producto_inventario[2]
        nueva_cantidad = cantidad_actual + cantidad_recibida

        inventario.actualizar_cantidad(cod_producto, nueva_cantidad)

    else:
        producto = detalle_producto.obtener_por_codigo(cod_producto)
        nombre_producto = producto[1]
        cantidad = cantidad_recibida
        fecha_ingreso_material = fecha_entrada
        tipo_movimiento = "entrada"
        stock_minimo = 10

        inventario.insertar_inventario(nombre_producto, cantidad, fecha_ingreso_material, tipo_movimiento, stock_minimo, id_orden, cod_producto)

    return jsonify({"mensaje": "Entrada de material registrada correctamente.",
                    "cumplimiento": cumplimiento_orden_compra,
                    "cantidad_recibida": cantidad_total,
                    "cantidad_pendiente": cantidad_ordenada - cantidad_total
                })


@app.route('/nueva_salida_material', methods=['POST'])
def crear_salida_material():

    cantidad_salida = request.json['cantidad_salida']
    cod_producto = request.json['cod_producto']

    producto_inventario = inventario.obtener_por_producto(cod_producto)

    if producto_inventario:
        cantidad_actual = producto_inventario[2]
        stock_minimo = producto_inventario[5]

        if  cantidad_actual - cantidad_salida >= stock_minimo:
            nueva_cantidad = cantidad_actual - cantidad_salida
            inventario.actualizar_cantidad(cod_producto, nueva_cantidad)
            return jsonify({"mensaje": "Salida del material correctamente"})
        else:
            return jsonify({"error": "No se puede realizar la salida porque el inventario quedaríapor debajo del stock mínimo."}), 400
    else:
        return jsonify({"error": "El producto no existe en el inventario."}), 404

@app.route('/empleados', methods=['GET'])
def listas_empleados():
    return jsonify(empleados.obtener_empleados())

@app.route('/empleado/<int:id>', methods=['GET'])
def listar_empleado(id):
    return jsonify(empleados.obtener_empleado(id))

@app.route('/nuevo_empleado', methods=['POST'])
def crear_empledo():
    ID_CC = request.json['ID_CC']
    Cargo_usuario = request.json['Cargo_usuario']
    Estado_Usuario = request.json['Estado_Usuario']
    fecha_ingreso = request.json['fecha_ingreso']
    area_asignada = request.json['area_asignada']
    id_usuario = request.json['id_usuario']

    empleados.insertar_empleado(ID_CC, Cargo_usuario, Estado_Usuario, fecha_ingreso, area_asignada, id_usuario)
    return jsonify({"mensaje": "Empleado creado correctamente."})

@app.route('/actualizar_empleado/<int:id>', methods=['PUT'])
def actualizar_empleado(id):
    data = request.json
    ID_CC = data.get('ID_CC')
    Cargo_usuario = data.get('Cargo_usuario')
    Estado_Usuario = data.get('Estado_Usuario')
    fecha_ingreso = data.get('fecha_ingreso')
    area_asignada = data.get('area_asignada')
    id_usuario = data.get('id_usuario')

    if ID_CC is None and Cargo_usuario is None and Estado_Usuario is None and fecha_ingreso is None and area_asignada is None and id_usuario is None:
        return jsonify({"error": "No se pudo actualizar los datos"})

    empleados_existente = empleados.obtener_empleado(id)

    if  empleados_existente:
        empleados.actualizar_empleado(id, ID_CC, Cargo_usuario, Estado_Usuario, fecha_ingreso, area_asignada, id_usuario)
        return jsonify({"mensaje": "El empleado esta actualizado correctamente"})
    else:
        return jsonify({"error": "El empleado no existe"}), 404

@app.route('/eliminar_empleado/<int:id>', methods=['DELETE'])
def eliminar_empleado(id):
    if empleados.obtener_empleado(id):
        empleados.eliminar_empleado(id)
        return jsonify({"mensaje": "Empleado eliminado correctamente."})
    else:
        return jsonify({"error": "El empleado no existe."}), 404

@app.route('/documentos_proveedor', methods=['GET'])
def listar_documentos_proveedor():
    return jsonify(documentos_proveedor.obtener_documentos_proveedor())

@app.route('/documento_proveedor/<int:id>', methods=['GET'])
def listar_documento_proveedor(id):
    return jsonify(documentos_proveedor.obtener_documento_proveedor(id))

@app.route('/ver_documento/<int:id>')
def ver_documento(id):
    documento = documentos_proveedor.obtener_documento_proveedor(id)
    if documento:
        ruta_archivo = documento[3]
        return send_file(ruta_archivo)
    return jsonify({"error": "El documento no existe"}), 404

@app.route('/nuevo_documento_proveedor', methods=['POST'])
def crear_documento_proveedor():
    tipo_documento = request.form['tipo_documento']
    archivo = request.files['archivo']
    id_proveedor = request.form['id_proveedor']

    nombre_archivo = archivo.filename
    fecha_subida = datetime.now().date()
    estado_validacion = 'Pendiente'
    ruta_carpeta = os.path.join(os.path.dirname(__file__), 'Documentos')
    ruta_archivo = os.path.join(ruta_carpeta, nombre_archivo)

    archivo.save(ruta_archivo)
    documentos_proveedor.insertar_documento_proveedor(tipo_documento, nombre_archivo, ruta_archivo, fecha_subida, estado_validacion, id_proveedor)
    return jsonify({"mensaje": "Docuementos cargados correctamente."})

@app.route('/actualizar_documento_proveedor/<int:id>', methods=['PUT'])
def actualizar_documento_proveedor(id):
    data = request.json
    tipo_documento = data.get('tipo_documento')
    nombre_archivo = data.get('nombre_archivo')
    ruta_archivo = data.get('ruta_archivo')
    fecha_subida = data.get('fecha_subida')
    estado_validacion = data.get('estado_validacion')
    id_proveedor = data.get('id_proveedor')

    if tipo_documento is None and nombre_archivo is None and ruta_archivo is None and fecha_subida is None and estado_validacion is None and id_proveedor is None:
        return jsonify({"error" : "No se pudo actualizar los datos."}), 400

    documento_proveedor_existente = documentos_proveedor.obtener_documento_proveedor(id)

    if documento_proveedor_existente:
        documentos_proveedor.actualizar_documento_proveedor(id, tipo_documento, nombre_archivo, ruta_archivo, fecha_subida, estado_validacion, id_proveedor)
        return jsonify({"mensaje": "Documento del proveedor actualizados correctamente."})
    else:
        return jsonify({"error": "El documento del proveedor no existe."}), 404

@app.route('/actualizar_estado_documento/<int:id>', methods=['PUT'])
def actualizar_estado_documento(id):
    data = request.json
    estado_validacion = data.get('estado_validacion')

    documento_existente = documentos_proveedor.obtener_documento_proveedor(id)

    if documento_existente:
        documentos_proveedor.actualizar_estado_documento(id, estado_validacion)
        return jsonify({"mensaje": "Estado del documento actualizado correctamente"})
    else:
        return jsonify({"error": "El documento no existe"}), 404

@app.route('/eliminar_documento_proveedor/<int:id>', methods=['DELETE'])
def eliminar_documento_proveedor(id):
    if documentos_proveedor.obtener_documento_proveedor(id):
        documentos_proveedor.eliminar_documento_proveedor(id)
        return jsonify({"mensaje": " Documento del proveedor eliminado correctamente."})
    else:
        return jsonify({"error": "El docuemento del proveedor no existe."}), 404

@app.route('/detalle_productos', methods=['GET'])
def listar_detalle_productos():
    return jsonify(detalle_producto.obtener_detalle_productos())

@app.route('/detalle_producto/<int:id>', methods=['GET'])
def listar_detalle_producto(id):
    return jsonify(detalle_producto.obtener_por_id(id))

@app.route('/detalle_producto_codigo/<int:cod_producto>', methods=['GET'])
def listar_producto_por_codigo(cod_producto):
    producto = detalle_producto.obtener_detalle_producto(cod_producto)

    if producto:
        return jsonify(producto)
    else:
        return jsonify({"error": "El producto no existe."}), 404

@app.route('/desactivar_detalle_producto/<int:id>', methods=['PUT'])
def desactivar_detalle_producto(id):
    detalle_producto_existente = detalle_producto.obtener_por_id(id)

    if detalle_producto_existente:
        detalle_producto.desactivar_detalle_producto(id)

        return jsonify({"mensaje": "El producto fue desactivado correctamente."})
    else:
        return jsonify({"error": "El producto no existe."}), 404

@app.route('/nuevo_detalle_producto', methods=['POST'])
def crear_detalle_producto():
    nombre_producto = request.json['nombre_producto']
    descripcion = request.json['descripcion']
    unidad_medida = request.json['unidad_medida']
    precio_unitario = request.json['precio_unitario']
    estado_producto = request.json['estado_producto']
    categoria =  request.json['categoria']
    marca =  request.json['marca']
    cod_producto = request.json['cod_producto']


    detalle_producto.insertar_detalle_producto(nombre_producto, descripcion, unidad_medida, precio_unitario, estado_producto, categoria, marca, cod_producto)
    return jsonify({"mensaje": "Detalle de producto creado correctamente."})

@app.route('/actualizar_detalle_producto/<int:id>', methods=['PUT'])
def actualizar_detalle_producto(id):
    data = request.json
    nombre_producto = data.get('nombre_producto')
    descripcion = data.get('descripcion')
    unidad_medida =  data.get('unidad_medida')
    precio_unitario = data.get('precio_unitario')
    estado_producto = data.get('estado_producto')
    categoria = data.get('categoria')
    marca = data.get('marca')
    cod_producto = data.get('cod_producto')


    if  nombre_producto is None and descripcion is None and unidad_medida is None and precio_unitario is None and estado_producto is None and categoria is None and marca is None and cod_producto is None:
        return jsonify({"error": "No se puede actualizar datos."}), 400

    detalle_producto_existente = detalle_producto.obtener_por_id(id)

    if detalle_producto_existente:
        detalle_producto.actualizar_detalle_producto(id, nombre_producto, descripcion, unidad_medida, precio_unitario, estado_producto, categoria, marca, cod_producto)
        return jsonify({"mensaje": "Detalle del producto se actualizo correctamente."})
    else:
        return jsonify({"error": "El detalle del producto no existe."}), 404

@app.route('/eliminar_detalle_producto/<int:id>', methods=['DELETE'])
def eliminar_detalle_producto(id):
    if detalle_producto.obtener_detalle_producto(id):
        detalle_producto.eliminar_detalle_producto(id)
        return jsonify({"mensaje": "Detalle del producto fue eliminado correctamente."})
    else:
        return jsonify({"error": "El detalle del producto no existe"}), 404

@app.route('/crear_pedidos', methods=['GET'])
def listar_crear_pedidos():
    return jsonify(crear_pedido.obtener_crear_pedidos())

@app.route('/pedidos_logistica', methods=['GET'])
def listar_pedidos_logistica():
    return jsonify(crear_pedido.obtener_pedidos_logistica())

@app.route('/crear_pedido/<int:id>', methods=['GET'])
def listar_crear_pedido(id):
    return jsonify(crear_pedido.obtener_crear_pedido(id))

@app.route('/nuevo_crear_pedido', methods=['POST'])
def crear_crear_pedido():
    fecha_creacion = request.json['fecha_creacion']
    nombre_producto = request.json['nombre_producto']
    cantidad = request.json['cantidad']
    destino = request.json['destino']
    fecha_requerida = request.json['fecha_requerida']
    estado_pedido = request.json['estado_pedido']
    id_usuario = request.json['id_usuario']
    cod_producto = request.json['cod_producto']

    id_pedido = crear_pedido.insertar_crear_pedido(fecha_creacion, nombre_producto, cantidad, destino, fecha_requerida, estado_pedido, id_usuario, cod_producto)
    historial_estado.insertar_historial_estado(estado_pedido, id_pedido)
    return jsonify({"mensaje": "Pedido creado correctamente.", "id_pedido": id_pedido})

@app.route('/actualizar_crear_pedido/<int:id>', methods=['PUT'])
def actualizar_crear_pedido(id):
    data = request.json
    fecha_creacion = data.get('fecha_creacion')
    nombre_producto = data.get('nombre_producto')
    cantidad = data.get('cantidad')
    destino = data.get('destino')
    fecha_requerida = data.get('fecha_requerida')
    estado_pedido = data.get('estado_pedido')
    id_usuario = data.get('id_usuario')
    cod_producto =  data.get('cod_producto')

    if fecha_creacion is None and  nombre_producto is None and cantidad is None and destino is None and fecha_requerida is None and estado_pedido is None and id_usuario is None and cod_producto is None:
        return jsonify({"error": "No se pueden actualizar los datos."}), 400

    crear_pedido_existente = crear_pedido.obtener_crear_pedido(id)

    if crear_pedido_existente:
        crear_pedido.actualizar_crear_pedido(id, fecha_creacion, nombre_producto, cantidad, destino, fecha_requerida, estado_pedido, id_usuario, cod_producto)
        return jsonify({"mensaje": "Pedido actualizado correctamente."})
    
    else:
        return jsonify({"error": "Crear pedido no existe."}), 404

@app.route('/eliminar_crear_pedido/<int:id>', methods=['DELETE'])
def eliminar_crear_pedido(id):
    if crear_pedido.obtener_crear_pedido(id):
        crear_pedido.eliminar_crear_pedido(id)
        return jsonify({"mensaje": "Crear pedido fue eliminado correctamente."})
    else:
        return jsonify({"error": "Crear pedido no existe."}), 404

@app.route('/reporte_pedidos', methods=['GET'])
def reporte_pedidos():
    resultado = crear_pedido.reporte_pedidos()
    return jsonify(resultado)

@app.route('/reporte_pedidos_excel', methods=['GET'])
def reporte_pedidos_execel():
    datos = crear_pedido.reporte_pedidos()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pedidos"

    encabezados =[
        "ID_pedido",
        "Usuario",
        "Fecha creación",
        "Nombre producto",
        "Cantidad",
        "Destino",
        "Fecha requerida",
        "Estado del pedido",
        "Código del producto",
        "Descripción",
        "Marca"
    ]
    hoja.append(encabezados)

    for fila in datos:
        hoja.append(fila)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)

    return send_file(
        archivo, as_attachment=True,
        download_name="reporte_pedidos.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route('/consultar_estado/<int:id_pedido>', methods=['GET'])
def listar_consultar_estado(id_pedido):
    resultado = consultar_estado.consultar_estado(id_pedido)

    print("Id pedido recibido:", id_pedido)
    print("resultado consulta:", resultado)

    if resultado:
        return jsonify(resultado)
    else:
        return jsonify({"error": "El pedido no existe."}),404

@app.route('/historial_estado/<int:id_pedido>', methods=['GET'])
def listar_histrial_estado(id_pedido):

    resultado = historial_estado.obtener_historial_estado(id_pedido)
    if resultado:
        return jsonify(resultado)
    else:
        return jsonify({"error": "No existe historial para este pedido"}), 404

@app.route('/cambiar_estado/<int:id_pedido>', methods=['PUT'])
def cambiar_estado(id_pedido):
    estado = request.json['estado']

    respuesta = historial_estado.cambiar_estado(id_pedido, estado)
    if respuesta:
        return jsonify({"mensaje": "Estado del pedido actualizado correctamente."})
    else:
        return jsonify({"error": "El pedido no existe."}), 404

if __name__ == '__main__':
    app.run(debug=True)
    









